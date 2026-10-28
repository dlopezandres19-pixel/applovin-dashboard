"""
AppLovin BuiltWith tracker
==========================

Pulls the BuiltWith Trends API for one technology (default: AppLovin), keeps a
weekly history, recomputes metrics, detects restatements and builds a
self-contained HTML dashboard.

Commands
--------
  python tracker.py backfill [--start 2025-01-04] [--dry-run]
      One-off: fills the weekly history back to --start (1 API credit per week).
      Weeks already stored are skipped, so re-running costs nothing extra.

  python tracker.py weekly [--recheck 6]
      Normal weekly run: 1 "current" pull (the only one that returns Expired)
      + 1 dated pull for this week + re-pulls the previous N weeks. The
      re-pulls pick up weeks BuiltWith publishes late (~2-3 week lag) and
      detect restatements. Default: 8 calls per run.

  python tracker.py build
      Recomputes metrics and writes docs/index.html. No API calls.

The API key is read from the BUILTWITH_API_KEY environment variable
(GitHub secret in Actions, or a local .env file).
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys
import time
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import requests

# ── configuration ────────────────────────────────────────────────────────────
TECH = os.environ.get("BUILTWITH_TECH", "AppLovin")
API_URL = "https://api.builtwith.com/trends/v6/api.json"
ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
SNAPSHOTS = DATA / "snapshots.csv"      # every API pull, never overwritten
WEEKLY = DATA / "weekly.csv"            # one row per week, with metrics
RESTATEMENTS = DATA / "restatements.csv"
CHURN = DATA / "churn.csv"              # gross adds / churn, one row per weekly run
PRICES = DATA / "app_prices.csv"        # APP daily close (Yahoo / Nasdaq / Stooq)
VALUATION_XLSX = DATA / "APP_Dashboard.xlsx"   # Capital IQ export, updated manually
VALUATION_TABS = {"pe": "PE NTM", "ev": "EVEBITDA NTM"}
DASHBOARD = ROOT / "docs" / "index.html"
TEMPLATE = ROOT / "dashboard_template.html"

ANCHOR_WEEKDAY = 5          # Saturday (BuiltWith chart points fall on Saturdays)
DEFAULT_START = "2025-01-04"
REQUEST_PAUSE_S = 1.0       # be polite with the API

SNAP_FIELDS = ["pulled_at_utc", "mode", "as_of", "ten_k", "hundred_k",
               "top_1m", "live", "live_plus_expired"]
RESTATE_FIELDS = ["detected_at_utc", "as_of", "field", "old_value",
                  "new_value", "diff"]
COUNT_FIELDS = ["live", "top_1m", "hundred_k", "ten_k"]


# ── helpers ──────────────────────────────────────────────────────────────────
def load_env_file() -> None:
    """Minimal .env loader so the script also runs locally."""
    env = ROOT / ".env"
    if not env.exists():
        return
    for line in env.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


def api_key() -> str:
    load_env_file()
    key = os.environ.get("BUILTWITH_API_KEY", "").strip()
    if not key:
        sys.exit("ERROR: falta BUILTWITH_API_KEY (secret de GitHub o archivo .env).")
    return key


def now_utc() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def last_anchor(d: date) -> date:
    """Most recent Saturday on or before d."""
    return d - timedelta(days=(d.weekday() - ANCHOR_WEEKDAY) % 7)


def anchors_between(start: date, end: date) -> list[date]:
    first = start + timedelta(days=(ANCHOR_WEEKDAY - start.weekday()) % 7)
    out, d = [], first
    while d <= end:
        out.append(d)
        d += timedelta(days=7)
    return out


def to_int(v) -> int | None:
    try:
        n = int(float(v))
    except (TypeError, ValueError):
        return None
    return None if n < 0 else n   # API returns -1 when a field is unavailable


def read_csv(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def append_csv(path: Path, fields: list[str], rows: list[dict]) -> None:
    if not rows:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    new = not path.exists()
    with path.open("a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        if new:
            w.writeheader()
        w.writerows(rows)


def write_csv(path: Path, fields: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


# ── API ──────────────────────────────────────────────────────────────────────
def fetch(key: str, as_of: date | None) -> dict:
    """One API call = 1 credit. as_of=None -> current totals (includes expired)."""
    params = {"KEY": key, "TECH": TECH}
    if as_of:
        params["DATE"] = as_of.isoformat()
    for attempt in range(3):
        try:
            r = requests.get(API_URL, params=params, timeout=60)
            r.raise_for_status()
            payload = r.json()
            break
        except (requests.RequestException, ValueError) as e:
            if attempt == 2:
                raise SystemExit(f"ERROR llamando a BuiltWith ({as_of}): {e}")
            time.sleep(5 * (attempt + 1))
    tech = payload.get("Tech") or payload.get("tech")
    if not tech or "coverage" not in tech:
        msg = json.dumps(payload)[:300]
        raise SystemExit(f"ERROR: respuesta inesperada de BuiltWith: {msg}")
    cov = tech["coverage"]
    return {
        "ten_k": to_int(cov.get("ten_k")),
        "hundred_k": to_int(cov.get("hundred_k")),
        "top_1m": to_int(cov.get("milly")),
        "live": to_int(cov.get("live")),
        # NB: the API's "expired" field is Live + Expired (cumulative ever seen)
        "live_plus_expired": to_int(cov.get("expired")),
    }


def snapshot_row(mode: str, as_of: date, values: dict) -> dict:
    row = {"pulled_at_utc": now_utc(), "mode": mode, "as_of": as_of.isoformat()}
    row.update({k: ("" if values[k] is None else values[k]) for k in values})
    return row


# ── restatement detection ────────────────────────────────────────────────────
def counts_key(s: dict) -> tuple:
    return tuple(to_int(s.get(f)) for f in COUNT_FIELDS)


def pending_weeks(dated: dict[str, dict]) -> set[str]:
    """BuiltWith publishes with a lag of ~2-3 weeks. Asking for a date past its
    latest data point returns that latest point again, so the newest weeks come
    back as exact copies of the last published week. Those trailing copies are
    'pending' (not yet published) and are left out of the series until real
    figures arrive on a later weekly run."""
    weeks = sorted(dated)
    pending: set[str] = set()
    for i in range(len(weeks) - 1, 0, -1):
        if counts_key(dated[weeks[i]]) == counts_key(dated[weeks[i - 1]]):
            pending.add(weeks[i])
        else:
            break
    return pending


def find_restatements(old_snaps: list[dict], new_rows: list[dict]) -> list[dict]:
    """Compare each new dated pull with the latest earlier dated pull for the
    same week. Current-mode pulls are never compared (different query type).
    Weeks that were still pending (copies of the prior week) are not
    restatements when real figures arrive - that is just publication."""
    latest: dict[str, dict] = {}
    for s in old_snaps:
        if s["mode"] == "date":
            latest[s["as_of"]] = s
    was_pending = pending_weeks(latest)
    out = []
    for n in new_rows:
        if n["mode"] != "date" or n["as_of"] not in latest or n["as_of"] in was_pending:
            continue
        prev = latest[n["as_of"]]
        for f in COUNT_FIELDS:
            a, b = to_int(prev.get(f)), to_int(n.get(f))
            if a is not None and b is not None and a != b:
                out.append({"detected_at_utc": n["pulled_at_utc"],
                            "as_of": n["as_of"], "field": f,
                            "old_value": a, "new_value": b, "diff": b - a})
    return out


# ── metrics ──────────────────────────────────────────────────────────────────
WEEKLY_FIELDS = [
    "week", "live", "top_1m", "hundred_k", "ten_k", "outside_top_1m",
    "top_1m_share_pct", "net_adds", "wow_pct", "wow_4w_avg_pct",
    "wow_accel_pp", "app_close_usd", "restated",
]
CHURN_FIELDS = ["pull_week", "pulled_at_utc", "live", "expired", "live_plus_expired",
                "gross_adds", "churned", "net_adds", "churn_pct"]


def pct(a, b):
    return None if a is None or b in (None, 0) else round((a - b) / b * 100, 2)


def build_weekly(snaps: list[dict], restatements: list[dict], prices: list[tuple]) -> list[dict]:
    """One row per published week. Counts come from the latest dated pull for
    that week (consistent query type over time). Weeks BuiltWith has not
    published yet are excluded. app_close_usd = last APP close on or before
    the week date (Friday close for a Saturday week)."""
    dated: dict[str, dict] = {}
    for s in snaps:
        if s["mode"] == "date":
            dated[s["as_of"]] = s                      # later rows win
    restated_weeks = {r["as_of"] for r in restatements}
    weeks = sorted(set(dated) - pending_weeks(dated))

    rows: list[dict] = []
    pi = 0
    for w in weeks:
        src = dated[w]
        r = {"week": w}
        for f in COUNT_FIELDS:
            r[f] = to_int(src.get(f))
        while pi < len(prices) and prices[pi][0] <= w:
            pi += 1
        r["app_close_usd"] = prices[pi - 1][1] if pi and (date.fromisoformat(w) - date.fromisoformat(prices[pi - 1][0])).days <= 7 else None
        r["outside_top_1m"] = (r["live"] - r["top_1m"]) if None not in (r["live"], r["top_1m"]) else None
        r["top_1m_share_pct"] = round(r["top_1m"] / r["live"] * 100, 1) if r["live"] and r["top_1m"] is not None else None
        r["restated"] = "yes" if w in restated_weeks else ""
        rows.append(r)

    for i, r in enumerate(rows):
        p = rows[i - 1] if i else None
        r["net_adds"] = (r["live"] - p["live"]) if p and None not in (r["live"], p["live"]) else None
        r["wow_pct"] = pct(r["live"], p["live"]) if p else None
        last4 = [x["wow_pct"] for x in rows[max(0, i - 3): i + 1] if x["wow_pct"] is not None]
        r["wow_4w_avg_pct"] = round(sum(last4) / len(last4), 2) if len(last4) == 4 else None
        r["wow_accel_pp"] = round(r["wow_pct"] - p["wow_pct"], 2) if p and None not in (r["wow_pct"], p.get("wow_pct")) else None
    return [{k: ("" if r.get(k) is None else r.get(k)) for k in WEEKLY_FIELDS} for r in rows]


def build_churn(snaps: list[dict]) -> list[dict]:
    """Gross adds and churn come only from 'current' pulls - the only query
    where BuiltWith returns expired sites. One row per weekly run (pull_week =
    Saturday of the run; the latest pull that week wins). Between two runs:
      gross_adds = change in live+expired  (websites first detected ever)
      churned    = change in expired       (net of sites that came back)
      net_adds   = change in live          (= gross_adds - churned)
      churn_pct  = churned / live at the previous run"""
    pulls: dict[str, dict] = {}
    for s in snaps:
        if s["mode"] == "current":
            pulls[s["as_of"]] = s
    rows, prev = [], None
    for wk in sorted(pulls):
        s = pulls[wk]
        live, lpe = to_int(s.get("live")), to_int(s.get("live_plus_expired"))
        if live is None or lpe is None:
            continue
        r = {"pull_week": wk, "pulled_at_utc": s["pulled_at_utc"], "live": live,
             "expired": lpe - live, "live_plus_expired": lpe,
             "gross_adds": None, "churned": None, "net_adds": None, "churn_pct": None}
        if prev:
            r["gross_adds"] = lpe - prev["live_plus_expired"]
            r["churned"] = r["expired"] - prev["expired"]
            r["net_adds"] = live - prev["live"]
            r["churn_pct"] = round(r["churned"] / prev["live"] * 100, 2) if prev["live"] else None
        rows.append(r)
        prev = r
    return [{k: ("" if r.get(k) is None else r.get(k)) for k in CHURN_FIELDS} for r in rows]


# ── APP share price (free public endpoints, no key) ─────────────────────────
PRICE_TICKER = os.environ.get("PRICE_TICKER", "APP")
PRICE_START = "2024-12-01"
UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126 Safari/537.36",
      "Accept": "application/json,text/csv,*/*"}


def _prices_yahoo() -> list[tuple]:
    t0 = int(datetime(2024, 12, 1, tzinfo=timezone.utc).timestamp())
    r = requests.get(f"https://query1.finance.yahoo.com/v8/finance/chart/{PRICE_TICKER}",
                     params={"period1": t0, "period2": int(time.time()), "interval": "1d"}, headers=UA, timeout=30)
    r.raise_for_status()
    res = r.json()["chart"]["result"][0]
    closes = res["indicators"]["quote"][0]["close"]
    return [(datetime.fromtimestamp(ts, timezone.utc).strftime("%Y-%m-%d"), round(c, 2))
            for ts, c in zip(res["timestamp"], closes) if c is not None]


def _prices_nasdaq() -> list[tuple]:
    r = requests.get(f"https://api.nasdaq.com/api/quote/{PRICE_TICKER}/historical",
                     params={"assetclass": "stocks", "fromdate": PRICE_START,
                             "todate": date.today().isoformat(), "limit": 9999},
                     headers={**UA, "Origin": "https://www.nasdaq.com", "Referer": "https://www.nasdaq.com/"}, timeout=30)
    r.raise_for_status()
    out = []
    for row in r.json()["data"]["tradesTable"]["rows"]:
        m, d, y = row["date"].split("/")
        out.append((f"{y}-{m}-{d}", round(float(row["close"].replace("$", "").replace(",", "")), 2)))
    return sorted(out)


def _prices_stooq() -> list[tuple]:
    key = os.environ.get("STOOQ_API_KEY", "").strip()
    if not key:
        raise RuntimeError("STOOQ_API_KEY not set")
    r = requests.get("https://stooq.com/q/d/l/", params={"s": f"{PRICE_TICKER.lower()}.us", "i": "d", "apikey": key},
                     headers=UA, timeout=30)
    r.raise_for_status()
    rows = list(csv.DictReader(r.text.splitlines()))
    return [(x["Date"], round(float(x["Close"]), 2)) for x in rows if x.get("Date", "") >= PRICE_START]


def update_prices() -> None:
    """Refreshes data/app_prices.csv. Never fails the run: if every source is
    down, the previous file is kept and the dashboard uses it as is."""
    load_env_file()
    for name, fn in (("Yahoo", _prices_yahoo), ("Nasdaq", _prices_nasdaq), ("Stooq", _prices_stooq)):
        try:
            rows = fn()
            if len(rows) > 50:
                write_csv(PRICES, ["date", "close"], [{"date": d, "close": c} for d, c in rows])
                print(f"Precios {PRICE_TICKER}: {len(rows)} días desde {name} (último {rows[-1][0]}: {rows[-1][1]})")
                return
            print(f"  {name}: respuesta muy corta ({len(rows)} filas)")
        except Exception as e:                          # noqa: BLE001
            print(f"  {name}: no disponible ({str(e)[:120]})")
    print("AVISO: no se pudieron actualizar los precios; se usa el archivo anterior si existe.")


def read_prices() -> list[tuple]:
    return [(r["date"], float(r["close"])) for r in read_csv(PRICES) if r.get("close")]


# ── valuation (Capital IQ workbook) ──────────────────────────────────────────
def read_valuation() -> dict | None:
    """Reads the 'PE NTM' and 'EVEBITDA NTM' tabs. Layout per tab: a row whose
    column B is 'Pricing Date', company names to its right, one row per month
    below. 'NM' (not meaningful) and blanks become null. Uses the values cached
    by Excel, so save the workbook after Capital IQ refreshes before uploading."""
    if not VALUATION_XLSX.exists():
        return None
    import openpyxl
    wb = openpyxl.load_workbook(VALUATION_XLSX, data_only=True, read_only=True)
    out: dict = {}
    for key, tab in VALUATION_TABS.items():
        if tab not in wb.sheetnames:
            continue
        rows = list(wb[tab].iter_rows(values_only=True))
        hdr = next((i for i, r in enumerate(rows) if len(r) > 1 and str(r[1]).strip().lower() == "pricing date"), None)
        if hdr is None:
            continue
        names = [str(c).strip() for c in rows[hdr][2:] if c not in (None, "")]
        tickers = [str(c).strip() if c else "" for c in rows[hdr - 1][2:2 + len(names)]] if hdr else [""] * len(names)
        dates, series = [], {n: [] for n in names}
        for r in rows[hdr + 1:]:
            d = r[1] if len(r) > 1 else None
            if not isinstance(d, (datetime, date)):
                continue
            dates.append(d.strftime("%Y-%m-%d"))
            for j, n in enumerate(names):
                v = r[2 + j] if len(r) > 2 + j else None
                series[n].append(round(float(v), 2) if isinstance(v, (int, float)) else None)
        out[key] = {"tab": tab, "dates": dates, "names": names, "tickers": tickers, "series": series}
    wb.close()
    return out or None


# ── dashboard ────────────────────────────────────────────────────────────────
def build_dashboard(refresh_prices: bool = True) -> None:
    if refresh_prices:
        update_prices()
    snaps = read_csv(SNAPSHOTS)
    restatements = read_csv(RESTATEMENTS)
    weekly = build_weekly(snaps, restatements, read_prices())
    churn = build_churn(snaps)
    write_csv(WEEKLY, WEEKLY_FIELDS, weekly)
    write_csv(CHURN, CHURN_FIELDS, churn)

    dated = {s["as_of"]: s for s in snaps if s["mode"] == "date"}
    payload = {
        "tech": TECH,
        "generated_at": now_utc(),
        "pending_weeks": sorted(pending_weeks(dated)),
        "weekly": weekly,
        "churn": churn,
        "price_ticker": PRICE_TICKER,
        "restatements": restatements[-200:],
        "valuation": read_valuation(),
    }
    html = TEMPLATE.read_text(encoding="utf-8").replace(
        "/*__DATA__*/null", json.dumps(payload, ensure_ascii=False).replace("</", "<\\/"))
    DASHBOARD.parent.mkdir(parents=True, exist_ok=True)
    DASHBOARD.write_text(html, encoding="utf-8")
    print(f"OK  {len(weekly)} semanas -> {WEEKLY.relative_to(ROOT)} y {DASHBOARD.relative_to(ROOT)}")


# ── commands ─────────────────────────────────────────────────────────────────
def cmd_backfill(args) -> None:
    start = date.fromisoformat(args.start)
    end = last_anchor(date.today()) - timedelta(days=7)   # this week is the weekly job's
    have = {s["as_of"] for s in read_csv(SNAPSHOTS) if s["mode"] == "date"}
    todo = [d for d in anchors_between(start, end) if d.isoformat() not in have]
    print(f"Backfill {TECH}: {len(todo)} semanas por bajar ({len(todo)} créditos).")
    if args.dry_run or not todo:
        return
    key = api_key()
    for i, d in enumerate(todo, 1):
        row = snapshot_row("date", d, fetch(key, d))
        append_csv(SNAPSHOTS, SNAP_FIELDS, [row])          # save as we go
        print(f"  [{i}/{len(todo)}] {d}  live={row['live']}")
        time.sleep(REQUEST_PAUSE_S)
    build_dashboard()


def cmd_weekly(args) -> None:
    key = api_key()
    today_anchor = last_anchor(date.today())
    old = read_csv(SNAPSHOTS)
    new_rows = [snapshot_row("current", today_anchor, fetch(key, None))]
    for k in range(0, args.recheck + 1):                 # this week + N previous
        d = today_anchor - timedelta(days=7 * k)
        time.sleep(REQUEST_PAUSE_S)
        new_rows.append(snapshot_row("date", d, fetch(key, d)))
    found = find_restatements(old, new_rows)
    append_csv(SNAPSHOTS, SNAP_FIELDS, new_rows)
    append_csv(RESTATEMENTS, RESTATE_FIELDS, found)
    cur = new_rows[0]
    print(f"Semana {today_anchor}: live={cur['live']}  top1m={cur['top_1m']}  "
          f"créditos usados={len(new_rows)}  restatements={len(found)}")
    for r in found:
        print(f"  RESTATEMENT {r['as_of']} {r['field']}: {r['old_value']} -> {r['new_value']}")
    build_dashboard()


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("backfill")
    b.add_argument("--start", default=DEFAULT_START)
    b.add_argument("--dry-run", action="store_true")
    w = sub.add_parser("weekly")
    w.add_argument("--recheck", type=int, default=6)
    sub.add_parser("build")
    args = p.parse_args()
    {"backfill": cmd_backfill, "weekly": cmd_weekly, "build": lambda a: build_dashboard()}[args.cmd](args)


if __name__ == "__main__":
    main()
